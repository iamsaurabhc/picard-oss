from app.services.docx_agent import build_docx_suggestion


def test_build_docx_suggestion():
    payload = build_docx_suggestion(
        document_id="doc-1",
        find="ACME Corp",
        replace="NewCo Inc.",
        change_mode="tracked",
        rationale="Standard party rename",
    )
    assert payload["document_id"] == "doc-1"
    assert payload["find"] == "ACME Corp"
    assert payload["replace"] == "NewCo Inc."
    assert payload["change_mode"] == "tracked"
    assert payload["rationale"] == "Standard party rename"


def test_docx_suggest_endpoint(client, monkeypatch):
    import io

    from docx import Document

    def _docx_bytes() -> bytes:
        doc = Document()
        doc.add_paragraph("Hello from DOCX suggest test.")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    monkeypatch.setattr("app.services.ingestion._executor.submit", lambda *args, **kwargs: None)
    ws = client.post("/workspaces", json={"name": "DOCX Suggest"}).json()
    files = {
        "file": (
            "memo.docx",
            _docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    doc = client.post(f"/workspaces/{ws['id']}/documents", files=files).json()
    r = client.post(
        f"/documents/{doc['id']}/docx/suggest",
        json={"pattern": "Hello", "replacement": "Hi", "tracked": True},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["document_id"] == doc["id"]
    assert body["find"] == "Hello"
    assert body["replace"] == "Hi"
    assert body["change_mode"] == "tracked"
